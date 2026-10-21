from io import BytesIO

import pytest
from docx import Document

from pilot_pipeline import analyze_document_bytes


def synthetic_docx_bytes():
    doc = Document()
    doc.add_paragraph("1. PAYMENT")
    doc.add_paragraph("Buyer shall pay 90 percent upon shipment.")
    doc.add_paragraph("2. DELIVERY")
    doc.add_paragraph("FOB Bandar Abbas.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_pipeline_generates_structured_and_persian_report():
    result = analyze_document_bytes(synthetic_docx_bytes(), "docx", "buyer", "SYNTHETIC-002")
    assert result["status"] == "human_review_required"
    assert result["screening"]["clause_count"] >= 2
    assert "پرداخت" in result["report_markdown"]
    assert "SYNTHETIC-002" in result["report_markdown"]
    assert result["extracted_characters"] > 20


def test_pipeline_rejects_unknown_role():
    with pytest.raises(ValueError):
        analyze_document_bytes(synthetic_docx_bytes(), "docx", "admin")


def test_pipeline_rejects_invalid_document():
    with pytest.raises(ValueError):
        analyze_document_bytes(b"not-a-document", "pdf", "seller")
