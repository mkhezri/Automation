from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from document_extract import extract_document_text


def test_docx_extract():
    document = Document()
    document.add_paragraph("1. PAYMENT")
    document.add_paragraph("Buyer shall pay 90 percent against documents.")
    output = BytesIO()
    document.save(output)
    result = extract_document_text(output.getvalue(), "docx")
    assert "PAYMENT" in result
    assert "Buyer shall pay" in result


def test_empty_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError):
        extract_document_text(output.getvalue(), "pdf")


def test_invalid_format_rejected():
    with pytest.raises(ValueError):
        extract_document_text(b"invalid document", "docx")


def test_unsupported_extension_rejected():
    with pytest.raises(ValueError):
        extract_document_text(b"sample", "txt")
