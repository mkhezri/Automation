"""Offline extraction of text from in-memory DOCX and text-based PDF bytes.

Never processes macros, executes attachments, accesses network or writes files.
PDF scans/images require a separately approved OCR process (not provided).
"""
from __future__ import annotations

from io import BytesIO

MAX_BYTES = 8 * 1024 * 1024
MAX_CHARS = 200_000
SUPPORTED = {"docx", "pdf"}


def extract_document_text(data: bytes, file_type: str) -> str:
    if file_type not in SUPPORTED:
        raise ValueError("Only docx and text-based pdf are supported")
    if not isinstance(data, bytes) or not data or len(data) > MAX_BYTES:
        raise ValueError("Invalid or oversized document")
    if file_type == "docx":
        if not data.startswith(b"PK"):
            raise ValueError("Invalid DOCX header")
        from docx import Document
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF header")
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF not supported")
        if len(reader.pages) > 100:
            raise ValueError("Too many pages")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    if not text.strip():
        raise ValueError("No selectable text found; scanned documents are not supported")
    if len(text) > MAX_CHARS:
        raise ValueError("Extracted text exceeds limit")
    return text
